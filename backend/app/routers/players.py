import io
import re
from typing import Any

from fastapi import APIRouter, Depends, File, HTTPException, UploadFile
from fastapi.responses import StreamingResponse
from openpyxl import Workbook, load_workbook
from sqlalchemy import inspect, text
from sqlalchemy.orm import Session

from ..auth import get_current_user
from ..database import engine, get_db
from ..models import Player
from ..schemas import (
    PLAYER_STATUS_LABELS,
    PLAYER_STATUSES,
    PlayerCreate,
    PlayerImportResult,
    PlayerOut,
    PlayerUpdate,
)

router = APIRouter(
    prefix="/players",
    tags=["players"],
    dependencies=[Depends(get_current_user)],
)

HEADER_ALIASES = {
    "nom": "last_name",
    "lastname": "last_name",
    "last_name": "last_name",
    "name": "last_name",
    "prenom": "first_name",
    "prénom": "first_name",
    "firstname": "first_name",
    "first_name": "first_name",
    "postes": "positions",
    "poste": "positions",
    "positions": "positions",
    "licence": "license_number",
    "license": "license_number",
    "n licence": "license_number",
    "n° licence": "license_number",
    "no licence": "license_number",
    "numero de licence": "license_number",
    "numéro de licence": "license_number",
    "license_number": "license_number",
    "numero": "number",
    "numéro": "number",
    "n°": "number",
    "maillot": "number",
    "n maillot": "number",
    "n° maillot": "number",
    "numero maillot": "number",
    "numéro maillot": "number",
    "statut": "status",
    "status": "status",
}

POSITION_ALIASES: dict[str, int] = {
    "1": 1,
    "pilier": 1,
    "pilier g": 1,
    "pilier gauche": 1,
    "pilier gauche (1)": 1,
    "2": 2,
    "talonneur": 2,
    "3": 3,
    "pilier d": 3,
    "pilier droit": 3,
    "4": 4,
    "2e ligne": 4,
    "deuxième ligne": 4,
    "seconde ligne": 4,
    "5": 5,
    "6": 6,
    "3e ligne": 6,
    "troisième ligne": 6,
    "3e ligne aile": 6,
    "flanker": 6,
    "7": 7,
    "8": 8,
    "n°8": 8,
    "n8": 8,
    "numero 8": 8,
    "numéro 8": 8,
    "3e ligne centre": 8,
    "9": 9,
    "mêlée": 9,
    "melee": 9,
    "demi de mêlée": 9,
    "demi de melee": 9,
    "10": 10,
    "ouverture": 10,
    "demi d'ouverture": 10,
    "demi douverture": 10,
    "11": 11,
    "ailier": 11,
    "ailier g": 11,
    "ailier gauche": 11,
    "12": 12,
    "centre": 12,
    "13": 13,
    "14": 14,
    "ailier d": 14,
    "ailier droit": 14,
    "15": 15,
    "arrière": 15,
    "arriere": 15,
}


def ensure_player_columns() -> None:
    """Ajoute last_name / positions / status si la table existait déjà."""
    inspector = inspect(engine)
    if "players" not in inspector.get_table_names():
        return
    cols = {c["name"] for c in inspector.get_columns("players")}
    dialect = engine.dialect.name
    with engine.begin() as conn:
        if "last_name" not in cols:
            conn.execute(
                text("ALTER TABLE players ADD COLUMN last_name VARCHAR(100) DEFAULT ''")
            )
        if "positions" not in cols:
            if dialect == "sqlite":
                conn.execute(
                    text("ALTER TABLE players ADD COLUMN positions TEXT DEFAULT '[]'")
                )
            else:
                conn.execute(
                    text(
                        "ALTER TABLE players ADD COLUMN positions JSON "
                        "DEFAULT '[]'::json"
                    )
                )
        if "status" not in cols:
            conn.execute(
                text(
                    "ALTER TABLE players ADD COLUMN status VARCHAR(30) "
                    "DEFAULT 'titulaire'"
                )
            )
        # Normalise d’éventuelles valeurs NULL
        if dialect == "sqlite":
            conn.execute(
                text("UPDATE players SET last_name = '' WHERE last_name IS NULL")
            )
            conn.execute(
                text("UPDATE players SET positions = '[]' WHERE positions IS NULL")
            )
            conn.execute(
                text(
                    "UPDATE players SET status = 'titulaire' "
                    "WHERE status IS NULL OR status = ''"
                )
            )
        else:
            conn.execute(
                text("UPDATE players SET last_name = '' WHERE last_name IS NULL")
            )
            conn.execute(
                text("UPDATE players SET positions = '[]'::json WHERE positions IS NULL")
            )
            conn.execute(
                text(
                    "UPDATE players SET status = 'titulaire' "
                    "WHERE status IS NULL OR status = ''"
                )
            )


def _norm_header(value: Any) -> str:
    text_val = str(value or "").strip().lower()
    text_val = text_val.replace("é", "e").replace("è", "e").replace("ê", "e")
    text_val = re.sub(r"\s+", " ", text_val)
    return text_val


def _parse_positions(raw: Any) -> list[int]:
    if raw is None or str(raw).strip() == "":
        return []
    if isinstance(raw, (int, float)):
        pos = int(raw)
        return [pos] if 1 <= pos <= 15 else []

    parts = re.split(r"[,;/|]+", str(raw))
    result: list[int] = []
    for part in parts:
        key = _norm_header(part)
        if not key:
            continue
        if key in POSITION_ALIASES:
            pos = POSITION_ALIASES[key]
        elif key.isdigit():
            pos = int(key)
        else:
            # "pilier (1)" etc.
            m = re.search(r"(\d{1,2})", key)
            pos = int(m.group(1)) if m else None
            if pos is None and key in POSITION_ALIASES:
                pos = POSITION_ALIASES[key]
        if pos is not None and 1 <= pos <= 15 and pos not in result:
            result.append(pos)
    return result


def _cell_str(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _parse_status(raw: Any) -> str:
    key = _norm_header(raw)
    if not key:
        return "titulaire"
    key = key.replace("ç", "c")
    aliases = {
        "titulaire": "titulaire",
        "remplacant": "remplacant",
        "apprentissage": "apprentissage",
        "apprenti": "apprentissage",
        "app": "apprentissage",
        "tit": "titulaire",
        "rempl": "remplacant",
    }
    mapped = aliases.get(key, key)
    return mapped if mapped in PLAYER_STATUSES else "titulaire"


def _format_positions(positions: list | None) -> str:
    if not positions:
        return ""
    return ", ".join(str(p) for p in positions)


def _players_ordered(db: Session) -> list[Player]:
    return db.query(Player).order_by(Player.last_name, Player.first_name).all()


def _pdf_text(value: str) -> str:
    """Normalise pour Helvetica (latin-1)."""
    replacements = {
        "é": "e",
        "è": "e",
        "ê": "e",
        "ë": "e",
        "à": "a",
        "â": "a",
        "ä": "a",
        "ù": "u",
        "û": "u",
        "ü": "u",
        "ô": "o",
        "ö": "o",
        "î": "i",
        "ï": "i",
        "ç": "c",
        "É": "E",
        "È": "E",
        "Ê": "E",
        "À": "A",
        "Ç": "C",
        "°": "o",
        "–": "-",
        "—": "-",
        "’": "'",
        "‘": "'",
    }
    text_val = value or ""
    for src, dst in replacements.items():
        text_val = text_val.replace(src, dst)
    return text_val.encode("latin-1", errors="replace").decode("latin-1")


@router.get("", response_model=list[PlayerOut])
def list_players(db: Session = Depends(get_db)):
    return _players_ordered(db)


@router.post("", response_model=PlayerOut, status_code=201)
def create_player(payload: PlayerCreate, db: Session = Depends(get_db)):
    if db.query(Player).filter(Player.license_number == payload.license_number).first():
        raise HTTPException(status_code=400, detail="Numéro de licence déjà utilisé")
    player = Player(**payload.model_dump())
    db.add(player)
    db.commit()
    db.refresh(player)
    return player


@router.get("/template.xlsx")
def download_template():
    wb = Workbook()
    ws = wb.active
    ws.title = "Joueurs"
    ws.append(["Nom", "Prénom", "Postes", "Numéro de licence", "Statut", "N° maillot"])
    ws.append(["Dupont", "Jean", "9, 10", "LIC-001", "Titulaire", 9])
    ws.append(["Martin", "Paul", "1;3", "LIC-002", "Remplaçant", 1])
    ws.append(["Bernard", "Luc", "talonneur, pilier", "LIC-003", "Apprentissage", 3])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=modele_joueurs.xlsx"},
    )


@router.get("/export.xlsx")
def export_players_xlsx(db: Session = Depends(get_db)):
    """Export au même format que le modèle : ré-importable via « Importer Excel »."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Joueurs"
    ws.append(["Nom", "Prénom", "Postes", "Numéro de licence", "Statut", "N° maillot"])
    for p in _players_ordered(db):
        ws.append(
            [
                p.last_name or "",
                p.first_name or "",
                _format_positions(p.positions),
                p.license_number or "",
                PLAYER_STATUS_LABELS.get(
                    getattr(p, "status", None) or "titulaire", "Titulaire"
                ),
                p.number or 0,
            ]
        )
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=liste_joueurs.xlsx"},
    )


@router.get("/export.pdf")
def export_players_pdf(db: Session = Depends(get_db)):
    from fpdf import FPDF

    players = _players_ordered(db)
    pdf = FPDF(orientation="L", unit="mm", format="A4")
    pdf.set_auto_page_break(auto=True, margin=12)
    pdf.add_page()
    pdf.set_font("Helvetica", "B", 16)
    pdf.cell(0, 10, "Liste des joueurs", new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", "", 10)
    pdf.cell(0, 6, f"{len(players)} joueur(s)", new_x="LMARGIN", new_y="NEXT")
    pdf.ln(4)

    headers = ["Nom", "Prenom", "Poste", "N licence", "Statut"]
    widths = [55, 45, 70, 55, 45]
    pdf.set_font("Helvetica", "B", 10)
    for header, width in zip(headers, widths):
        pdf.cell(width, 8, header, border=1)
    pdf.ln()

    pdf.set_font("Helvetica", "", 9)
    for p in players:
        row = [
            _pdf_text(p.last_name or ""),
            _pdf_text(p.first_name or ""),
            _pdf_text(_format_positions(p.positions) or "-"),
            _pdf_text(p.license_number or ""),
            _pdf_text(
                PLAYER_STATUS_LABELS.get(
                    getattr(p, "status", None) or "titulaire", "Titulaire"
                )
            ),
        ]
        for value, width in zip(row, widths):
            pdf.cell(width, 7, value[:40], border=1)
        pdf.ln()

    out = pdf.output()
    buf = io.BytesIO(
        out if isinstance(out, (bytes, bytearray)) else out.encode("latin-1")
    )
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/pdf",
        headers={"Content-Disposition": "attachment; filename=liste_joueurs.pdf"},
    )


@router.post("/import", response_model=PlayerImportResult)
async def import_players(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
):
    filename = (file.filename or "").lower()
    if not filename.endswith((".xlsx", ".xlsm")):
        raise HTTPException(
            status_code=400,
            detail="Format non supporté. Utilisez un fichier Excel (.xlsx).",
        )

    content = await file.read()
    try:
        wb = load_workbook(io.BytesIO(content), data_only=True)
    except Exception as exc:  # noqa: BLE001
        raise HTTPException(status_code=400, detail=f"Fichier Excel invalide: {exc}") from exc

    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        raise HTTPException(status_code=400, detail="Fichier vide")

    header_map: dict[int, str] = {}
    for idx, cell in enumerate(rows[0]):
        key = HEADER_ALIASES.get(_norm_header(cell))
        if key:
            header_map[idx] = key

    required = {"last_name", "first_name", "license_number"}
    present_fields = set(header_map.values())
    if not required.issubset(present_fields):
        raise HTTPException(
            status_code=400,
            detail="Colonnes requises : Nom, Prénom, Numéro de licence (Postes / Statut / N° maillot optionnels).",
        )

    created = updated = skipped = 0
    errors: list[str] = []

    for row_idx, row in enumerate(rows[1:], start=2):
        if not row or all(c is None or str(c).strip() == "" for c in row):
            continue

        data: dict[str, Any] = {
            "number": 0,
            "first_name": "",
            "last_name": "",
            "license_number": "",
            "positions": [],
            "status": "titulaire",
        }
        for col_idx, field in header_map.items():
            if col_idx >= len(row):
                continue
            raw = row[col_idx]
            if field == "positions":
                data["positions"] = _parse_positions(raw)
            elif field == "status":
                data["status"] = _parse_status(raw)
            elif field == "number":
                try:
                    data["number"] = int(float(raw)) if raw not in (None, "") else 0
                except (TypeError, ValueError):
                    data["number"] = 0
            else:
                data[field] = _cell_str(raw)

        if not data["first_name"] or not data["last_name"] or not data["license_number"]:
            skipped += 1
            errors.append(f"Ligne {row_idx}: nom, prénom et licence obligatoires")
            continue

        existing = (
            db.query(Player)
            .filter(Player.license_number == data["license_number"])
            .first()
        )
        try:
            if existing:
                # Mise à jour par n° de licence (ré-import après export)
                existing.first_name = data["first_name"]
                existing.last_name = data["last_name"]
                if "positions" in present_fields:
                    existing.positions = data["positions"]
                if "status" in present_fields:
                    existing.status = data["status"]
                if "number" in present_fields:
                    existing.number = data["number"]
                updated += 1
            else:
                db.add(Player(**data))
                created += 1
        except Exception as exc:  # noqa: BLE001
            skipped += 1
            errors.append(f"Ligne {row_idx}: {exc}")

    db.commit()
    return PlayerImportResult(
        created=created,
        updated=updated,
        skipped=skipped,
        errors=errors[:50],
    )


@router.get("/{player_id}", response_model=PlayerOut)
def get_player(player_id: int, db: Session = Depends(get_db)):
    player = db.get(Player, player_id)
    if not player:
        raise HTTPException(status_code=404, detail="Joueur introuvable")
    return player


@router.put("/{player_id}", response_model=PlayerOut)
def update_player(player_id: int, payload: PlayerUpdate, db: Session = Depends(get_db)):
    player = db.get(Player, player_id)
    if not player:
        raise HTTPException(status_code=404, detail="Joueur introuvable")
    data = payload.model_dump(exclude_unset=True)
    if "license_number" in data:
        existing = (
            db.query(Player)
            .filter(Player.license_number == data["license_number"], Player.id != player_id)
            .first()
        )
        if existing:
            raise HTTPException(status_code=400, detail="Numéro de licence déjà utilisé")
    for key, value in data.items():
        setattr(player, key, value)
    db.commit()
    db.refresh(player)
    return player


@router.delete("/{player_id}", status_code=204)
def delete_player(player_id: int, db: Session = Depends(get_db)):
    player = db.get(Player, player_id)
    if not player:
        raise HTTPException(status_code=404, detail="Joueur introuvable")
    db.delete(player)
    db.commit()
